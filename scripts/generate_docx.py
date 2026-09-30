"""
Generate PROJECT_SYNOPSIS.docx - Compact format matching college reference
Department: Department of Computer Science(IoT Cybersecurity Including Blockchain)
"""
import sys
import io
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

DEPT = 'Department of Computer Science(IoT Cybersecurity Including Blockchain)'
COLLEGE = "Alva's Institute of Engineering & Technology, Moodbidri"


def set_run(run, size=12, bold=False, font_name='Times New Roman'):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = font_name


def add_centered(doc, text, size=12, bold=False, after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(14)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p


def add_body(doc, text, size=12, bold=False, after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(15)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p


def add_heading(doc, text, size=12, after=4, new_page=False):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(15)
    r = p.add_run(text)
    set_run(r, size=size, bold=True)
    return p


def add_bullet(doc, title, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(15)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run('- ')
    set_run(r, size=size, bold=True)
    r2 = p.add_run(f'{title}: ')
    set_run(r2, size=size, bold=True)
    r3 = p.add_run(text)
    set_run(r3, size=size)
    return p


def add_simple_bullet(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(15)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run('- ')
    set_run(r, size=size)
    r2 = p.add_run(text)
    set_run(r2, size=size)
    return p


def add_lit_bullet(doc, author_text, description, source_label, source_url=None, size=11):
    """Add a literature survey bullet point: bold author, regular description, underlined source."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(14)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    # Bullet symbol
    r_bullet = p.add_run('\u2022 ')
    set_run(r_bullet, size=size, bold=True)
    # Bold author citation
    r_author = p.add_run(author_text + ' ')
    set_run(r_author, size=size, bold=True)
    # Regular description
    r_desc = p.add_run(description + ' ')
    set_run(r_desc, size=size)
    # Underlined source
    r_src = p.add_run(source_label)
    set_run(r_src, size=size)
    r_src.underline = True
    from docx.shared import RGBColor as _RGBColor
    r_src.font.color.rgb = _RGBColor(0x00, 0x00, 0xFF)
    return p


def generate_docx():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ==================== PAGE 1: TITLE PAGE ====================
    # Add VTU header image at the top
    import os
    from docx.shared import RGBColor
    from docx.oxml.ns import qn

    vtu_img = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'images', 'vtu_header.png')
    if os.path.exists(vtu_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run()
        run.add_picture(vtu_img, width=Inches(5.5))

    # Title
    add_centered(doc, 'PROPOSED PROJECT SYNOPSIS', size=14, bold=True, after=10)
    add_centered(doc, 'by', size=12, bold=True, after=10)

    # Team members table — name LEFT, USN RIGHT with underlines
    members = [
        ('Arya B Shetty', '4AL23IC007'),
        ('Harinand J', '4AL23IC012'),
        ('Karthik P', '4AL23IC016'),
        ('Pavan SN', '4AL23IC033'),
    ]

    table = doc.add_table(rows=len(members), cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (name, usn) in enumerate(members):
        # Name cell - left aligned, underlined
        cell_name = table.cell(i, 0)
        cell_name.width = Inches(3)
        p = cell_name.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(name)
        set_run(r, size=12)
        r.underline = True

        # USN cell - right aligned, underlined
        cell_usn = table.cell(i, 1)
        cell_usn.width = Inches(3)
        p = cell_usn.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(usn)
        set_run(r, size=12)
        r.underline = True

    # Remove table borders
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = tcPr.find(qn('w:tcBorders'))
            if tcBorders is not None:
                tcPr.remove(tcBorders)

    # Remove table-level borders
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        borders = tblPr.find(qn('w:tblBorders'))
        if borders is not None:
            tblPr.remove(borders)

    add_centered(doc, '', size=8, after=10)

    # Department and College — underlined, centered
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(DEPT)
    set_run(r, size=12)
    r.underline = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(COLLEGE)
    set_run(r, size=12)
    r.underline = True

    # Guide section
    add_centered(doc, 'Under the Guidance of', size=12, bold=True, after=8)

    # Guide name — bold + italic
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('Prof. Jyothibha R Chinchankar')
    set_run(r, size=12, bold=True)
    r.italic = True

    # Guide designation
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(15)
    r = p.add_run('Senior Associate Professor, ' + DEPT)
    set_run(r, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(COLLEGE)
    set_run(r, size=12)

    # ==================== PAGE 2: TABLE OF CONTENTS ====================
    doc.add_page_break()

    add_centered(doc, 'TABLE OF CONTENTS', size=13, bold=True, after=12)

    toc_items = [
        '1.  TITLE OF THE PROJECT',
        '2.  INTRODUCTION',
        '3.  PROBLEM STATEMENT',
        '4.  LITERATURE SURVEY',
        '5.  OBJECTIVES OF THE PROPOSED PROJECT',
        '6.  PROPOSED METHODOLOGY',
        '7.  HARDWARE REQUIREMENTS',
        '8.  SOFTWARE REQUIREMENTS',
        '9.  EXPECTED OUTCOME OF THE PROPOSED PROJECT',
        '10. SUMMARY',
        '11. REFERENCES',
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(18)
        r = p.add_run(item)
        set_run(r, size=12, bold=True)

    # ==================== PAGE 3: TITLE + ABSTRACT + INTRO + PROBLEM ====================
    doc.add_page_break()

    add_heading(doc, 'TITLE', size=12, after=4)
    add_body(doc,
        'AI Smart Glasses using ESP32-CAM with Hybrid Cloud-Local AI Architecture',
        size=12, after=6)

    add_heading(doc, 'ABSTRACT', size=12, after=4)
    add_body(doc,
        'This project presents the design and implementation of AI-powered smart glasses '
        'built using the ESP32-CAM microcontroller. The system employs a hybrid cloud-local '
        'architecture where heavy AI tasks (image description via Llama 4 Scout 17B, Q&A '
        'and translation via Llama 3.3 70B) are processed through the Groq free-tier cloud '
        'API, while lightweight tasks (text extraction via EasyOCR, speech recognition via '
        'Whisper-tiny) run locally. A Flask-based Smart Router ensures only one local model '
        'is loaded at a time, keeping peak RAM usage under 600 MB. The hardware prototype '
        'costs under Rs.1,500, making it an affordable assistive tool for the visually impaired.',
        size=12, after=6)

    add_heading(doc, 'INTRODUCTION', size=12, after=4)
    add_body(doc,
        'Artificial Intelligence has made remarkable progress in understanding images, '
        'processing speech, and generating human-like text responses. At the same time, '
        'microcontrollers and camera modules have become small and affordable enough to be '
        'embedded into everyday wearable devices. This project brings these two trends '
        'together by building AI-powered smart glasses using the ESP32-CAM microcontroller.',
        size=12, after=4)

    add_body(doc,
        'The glasses capture images through the onboard OV2640 camera and send them over '
        'Wi-Fi to a Flask-based backend server running on a local PC. The backend uses a '
        'Smart Routing architecture to distribute tasks across multiple AI models. Heavy '
        'tasks like image description and question-answering are offloaded to cloud-based '
        'large language models (Llama 4 Scout 17B for vision, Llama 3.3 70B for language) '
        'accessed through the Groq API free tier, which provides 14,400 free API calls per '
        'day. Lightweight tasks like text extraction (OCR) and speech-to-text transcription '
        'are handled locally using EasyOCR and OpenAI Whisper-tiny respectively.',
        size=12, after=4)

    add_body(doc,
        'The key innovation is the Smart Routing system, which ensures only one local model '
        'is loaded into memory at a time, keeping peak RAM usage under 600 MB. This makes '
        'the system runnable on consumer hardware like the NVIDIA RTX 3050, without requiring '
        'expensive cloud servers or high-end GPUs. The AI response is displayed on a small '
        'OLED screen and can also be spoken aloud through a speaker using text-to-speech. '
        'The entire hardware prototype costs under Rs.1,500.',
        size=12, after=6)

    add_heading(doc, 'PROBLEM STATEMENT', size=12, after=4)
    add_body(doc,
        'Existing AI-powered wearable devices have three major limitations that prevent '
        'widespread adoption:',
        size=12, after=4)

    add_body(doc,
        'Commercial smart glasses like Google Glass (Rs.83,000) and Meta Ray-Ban (Rs.25,000) '
        'are far too expensive for the average user. Running modern AI vision models requires '
        'powerful GPUs with 12+ GB of VRAM, making it impractical on consumer hardware. '
        'Visually impaired users need real-time help understanding their surroundings - '
        'reading signboards, identifying objects, and getting answers to questions about what '
        'they see. Existing affordable devices like the white cane cannot provide any visual '
        'understanding or text reading capabilities.',
        size=12, after=4)

    add_body(doc,
        'This project addresses all three problems by building a smart glasses prototype for '
        'under Rs.1,500 using the ESP32-CAM, and implementing a hybrid cloud-local AI '
        'architecture that runs on a standard laptop with just 600 MB of RAM, using free '
        'cloud APIs for the heavy AI processing.',
        size=12, after=4)

    # ==================== LITERATURE SURVEY (after Problem Statement) ====================
    add_heading(doc, 'LITERATURE SURVEY', size=12, after=6, new_page=True)

    add_lit_bullet(doc,
        'Das et al. (2021)',
        'proposed a low-cost smart glass system using the ESP-32 microcontroller, demonstrating '
        'that affordable embedded platforms could serve as the foundation for wearable computing '
        'with real-time image capture and wireless data transmission.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'Ramesh Babu et al. (2025)',
        'developed AI-powered smart glasses for the blind using OpenCV and Python, implementing '
        'computer vision pipelines in a wearable device to provide descriptive scene feedback '
        'to visually impaired users.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'Lin et al. (2020)',
        'proposed a smart glasses application system for visually impaired people based on deep '
        'learning, using convolutional neural networks for object recognition with audio feedback '
        'to improve scene understanding accuracy.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'Radford et al. (2023)',
        'introduced OpenAI Whisper, a general-purpose speech recognition model trained on 680,000 '
        'hours of multilingual data, achieving robust performance across diverse accents and noise '
        'conditions with models ranging from 39M to 1.5B parameters.',
        'proceedings.mlr.press')

    add_lit_bullet(doc,
        'Touvron et al. (2023)',
        'presented Llama 2, an open-source family of large language models ranging from 7B to 70B '
        'parameters, demonstrating competitive performance with proprietary systems on question '
        'answering, summarization, and translation tasks.',
        'ai.meta.com')

    add_lit_bullet(doc,
        'JaidedAI (2020)',
        'developed EasyOCR, an open-source optical character recognition library supporting 80+ '
        'languages, employing a CRAFT text detection network and CRNN recognition pipeline with '
        'a modest memory footprint of approximately 180 MB RAM.',
        'github.com/JaidedAI/EasyOCR')

    add_lit_bullet(doc,
        'Gospodinov and Krastev (2022)',
        'integrated object detection algorithms with smart glasses hardware, demonstrating that '
        'computer vision models could be coupled with head-mounted devices to provide '
        'context-aware information to the wearer.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'Sharma et al. (2025)',
        'presented an AI-driven wearable system for autonomous navigation using smart glasses, '
        'highlighting the importance of multi-modal output combining audio and visual display for '
        'communicating AI-generated information to visually impaired users.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'U. R. et al. (2023)',
        'developed "Third Eye," a smart wearable glass powered by deep learning and artificial '
        'intelligence, combining object detection, text recognition, and depth estimation '
        'capabilities in a single wearable platform for visually impaired people.',
        'ieeexplore.ieee.org')

    add_lit_bullet(doc,
        'Chen and Ran (2019)',
        'provided a comprehensive review of deep learning with edge computing, advocating for '
        'intelligent task partitioning that executes lightweight models at the edge while '
        'offloading compute-intensive tasks to the cloud to balance performance and resource '
        'constraints.',
        'ieeexplore.ieee.org')

    # ==================== OBJECTIVES + METHODOLOGY ====================
    add_heading(doc, 'OBJECTIVES OF THE PROPOSED PROJECT', size=12, after=4, new_page=True)
    add_body(doc, 'The objectives of this proposed project are:', size=12, after=4)

    add_simple_bullet(doc,
        'To design a wearable smart glasses prototype using ESP32-CAM with OV2640 camera, '
        'OLED display, INMP441 microphone, and speaker at a total cost under Rs.1,500.')

    add_simple_bullet(doc,
        'To implement and integrate hybrid cloud-local AI architecture using Groq free-tier '
        'API for heavy processing and local models for lightweight tasks.')

    add_simple_bullet(doc,
        'To enable real-time scene description, text extraction (OCR), speech-to-text, '
        'question answering, and translation capabilities.')

    add_simple_bullet(doc,
        'To achieve response times under 2 seconds for all AI tasks with peak memory usage '
        'under 600 MB through Smart Routing.')

    add_simple_bullet(doc,
        'To provide an affordable and accessible assistive tool for visually impaired users '
        'in the Indian market.')

    add_heading(doc, 'PROPOSED METHODOLOGY', size=12, after=4)

    add_body(doc,
        'The system follows a four-stage pipeline:',
        size=12, after=4)

    add_heading(doc, 'Step 1: Image/Audio Capture', size=11, after=2)
    add_body(doc,
        'The user presses a button on the glasses to select a mode. The ESP32-CAM captures a '
        'JPEG image at 640x480 resolution using the OV2640 camera, or the INMP441 microphone '
        'records audio for speech-to-text mode.',
        size=12, after=4)

    add_heading(doc, 'Step 2: Data Transmission', size=11, after=2)
    add_body(doc,
        'The captured image or audio is encoded in Base64 format and sent as an HTTP POST '
        'request over Wi-Fi to the Flask backend server running on the local PC.',
        size=12, after=4)

    add_heading(doc, 'Step 3: Smart Routing & AI Processing', size=11, after=2)
    add_body(doc,
        'The Flask Smart Router examines the request type and routes it to the appropriate '
        'AI model: Image description goes to Llama 4 Scout via Groq API (cloud). Text '
        'extraction goes to EasyOCR (local, ~180 MB). Speech transcription goes to '
        'Whisper-tiny (local, ~145 MB). Q&A and translation go to Llama 3.3 70B via Groq '
        'API (cloud). After processing, local models are unloaded from memory.',
        size=12, after=4)

    add_heading(doc, 'Step 4: Output & Display', size=11, after=2)
    add_body(doc,
        'The AI-generated response is sent back to the ESP32 as JSON. The response text is '
        'displayed on the SSD1306 OLED screen (128x64 pixels) and spoken aloud through the '
        'connected speaker using text-to-speech conversion.',
        size=12, after=4)

    # ==================== PAGE 5: HARDWARE + SOFTWARE ====================
    add_heading(doc, 'HARDWARE REQUIREMENTS', size=12, after=4, new_page=True)

    hw_items = [
        ('ESP32-CAM with OV2640 Camera (Rs.450)',
         '32-bit microcontroller with built-in Wi-Fi, Bluetooth, and 2MP camera module.'),
        ('OLED Display SSD1306, 0.96", I2C (Rs.180)',
         '128x64 pixel display for showing AI responses. Connected via I2C with 4.7k pull-ups.'),
        ('INMP441 I2S Digital Microphone (Rs.250)',
         'High-precision MEMS microphone for speech-to-text via I2S protocol.'),
        ('Micro Speaker 8 ohm, 0.5W (Rs.99)',
         'Audio output for text-to-speech feedback to the user.'),
        ('Push Buttons x2 (Rs.20)',
         'Mode selection and action trigger with 10k pull-up resistors.'),
        ('LiPo Battery 3.7V 1000mAh + TP4056 (Rs.240)',
         'Rechargeable battery with USB charging module for portable operation.'),
        ('AMS1117-3.3V Regulator (Rs.18)',
         'Converts 3.7V battery to stable 3.3V for all components.'),
        ('Supporting Components (Rs.150)',
         'Resistors, capacitors, jumper wires, and glasses frame for mounting.'),
    ]
    for title, desc in hw_items:
        add_bullet(doc, title, desc, size=12)

    add_body(doc, 'Total estimated hardware cost: Rs.1,407 (approximately $17 USD)',
             size=12, bold=True, after=6)

    add_heading(doc, 'SOFTWARE REQUIREMENTS', size=12, after=4)

    add_heading(doc, 'Programming Languages & Frameworks:', size=11, after=2)
    add_body(doc,
        'Python 3.10+ for Flask backend and AI integration. C/C++ (Arduino) for ESP32-CAM '
        'firmware. Flask 3.0+ as the Smart Router web framework.',
        size=12, after=4)

    add_heading(doc, 'AI Models & Libraries:', size=11, after=2)
    add_body(doc,
        'Groq Python SDK (v1.2) for cloud API access to Llama 4 Scout 17B (vision) and '
        'Llama 3.3 70B Versatile (language). EasyOCR 1.7+ for local OCR supporting 80+ '
        'languages. OpenAI Whisper (tiny, 39M parameters, ~145 MB) for local speech-to-text. '
        'PyTorch 2.0+ as the deep learning inference framework.',
        size=12, after=4)

    add_heading(doc, 'Development Tools:', size=11, after=2)
    add_body(doc,
        'Visual Studio Code for Python development. Arduino IDE 2.x for ESP32 firmware. '
        'python-dotenv for secure API key management. psutil for system resource monitoring.',
        size=12, after=4)

    # ==================== PAGE 6: EXPECTED OUTCOME + SUMMARY ====================
    add_heading(doc, 'EXPECTED OUTCOME OF THE PROPOSED PROJECT', size=12, after=4, new_page=True)

    add_simple_bullet(doc,
        'A working smart glasses prototype capturing images and displaying AI descriptions.')
    add_simple_bullet(doc,
        'A Flask backend with Smart Routing distributing requests between cloud and local models.')
    add_simple_bullet(doc,
        'Five functional AI capabilities: scene description, OCR, speech-to-text, Q&A, translation.')
    add_simple_bullet(doc,
        'Verified response times: Vision 0.44s, Q&A 0.54s, Translation 0.44s, OCR 1.82s, '
        'Speech-to-Text 0.37s.')
    add_simple_bullet(doc,
        'Improved memory efficiency with peak usage under 600 MB using lazy loading of models.')
    add_simple_bullet(doc,
        'Zero operational cost using Groq free tier (14,400 requests/day) and open-source models.')
    add_simple_bullet(doc,
        'A cost-effective prototype under Rs.1,500 for manufacturers to mass produce an '
        'affordable assistive device.')

    add_heading(doc, 'SUMMARY', size=12, after=4)

    add_body(doc,
        'This project demonstrates that advanced AI capabilities - image understanding, text '
        'extraction, speech recognition, and question answering - can be delivered through an '
        'affordable wearable device costing under Rs.1,500. The core hardware is the ESP32-CAM '
        'microcontroller, which captures images via its OV2640 camera and sends them over Wi-Fi '
        'to a Python Flask backend.',
        size=12, after=4)

    add_body(doc,
        'The backend implements a Smart Routing architecture that intelligently distributes '
        'AI tasks: heavy vision and language processing is offloaded to Groq\'s free cloud API '
        '(Llama 4 Scout for image description, Llama 3.3 70B for Q&A and translation), while '
        'lightweight tasks like OCR and speech-to-text run locally using EasyOCR and Whisper-tiny. '
        'The Smart Routing system ensures that only one local AI model is loaded into memory at '
        'any given time, keeping peak RAM at approximately 600 MB.',
        size=12, after=4)

    add_body(doc,
        'All 13 tests passed: 5/5 AI model tests (Llama 4 Scout vision, Llama 3.3 Q&A, '
        'Llama 3.3 translation, EasyOCR, Whisper-tiny), 3/3 memory management tests, and '
        '5/5 Flask server endpoint tests. The system achieves response times under 2 seconds '
        'for every AI task while costing nothing to operate, making it a practical and scalable '
        'assistive technology solution.',
        size=12, after=4)

    # ==================== PAGE 7: REFERENCES ====================
    add_heading(doc, 'REFERENCES', size=12, after=6, new_page=True)

    references = [
        'S. Das, S. Saxena, and N. K. Rout, "Low Cost Smart-Glass using ESP-32," in 2021 IEEE '
        '2nd International Conference on Applied Electromagnetics, Signal Processing, & '
        'Communication (AESPC), 2021.',

        'S. S. V. Ramesh Babu, M. Gowtham, G. Jayakrishna, and A. L. S. S. Manasa, '
        '"AI-Powered Smart Glasses For The Blind Using OpenCV & Python," in 2025 IEEE '
        'Wireless Antenna and Microwave Symposium (WAMS), 2025.',

        'J. Y. Lin, C. C. Yao, C. L. Chiang, and M. C. Chen, "Smart Glasses Application '
        'System for Visually Impaired People Based on Deep Learning," in Indo-Taiwan 2nd '
        'International Conference on Computing, Analytics and Networks (Indo-Taiwan ICAN), 2020.',

        'A. Radford, J. W. Kim, T. Xu, G. Brockman, C. McLeavey, and I. Sutskever, "Robust '
        'Speech Recognition via Large-Scale Weak Supervision," in Proceedings of the 40th '
        'International Conference on Machine Learning (ICML), 2023.',

        'H. Touvron, L. Martin, K. Stone, et al., "Llama 2: Open Foundation and Fine-Tuned '
        'Chat Models," Meta AI Research, 2023.',

        'JaidedAI, "EasyOCR: Ready-to-Use OCR with 80+ Supported Languages," GitHub, 2020. '
        'Available: https://github.com/JaidedAI/EasyOCR',

        'N. Gospodinov and G. Krastev, "Object detection with smart glasses," in 2022 30th '
        'Telecommunications Forum (TELFOR), Belgrade, Serbia, Nov. 2022.',

        'P. Sharma, M. Sadasivan, A. S. Babu, and K. T. A. Robert, "Smart Glasses for the '
        'Blind: A Real-Time AI-Driven Wearable System for Autonomous Navigation," in 2025 '
        'International Conference on Computing and Communications (COMPUTINGCON), Pune, India, '
        'Sep. 2025.',

        'U. R, H. P.S., R. N.G., and N. A, "Third Eye - A Smart Wearable Glass With Deep '
        'Learning Technology Powered With Artificial Intelligence for Visually Impaired People," '
        'in 2023 IEEE International Students\' Conference on Electrical, Electronics and '
        'Computer Science (SCEECS), 2023.',

        'J. Chen and X. Ran, "Deep Learning with Edge Computing: A Review," in Proceedings '
        'of the IEEE, vol. 107, no. 8, pp. 1655-1674, 2019.',
    ]

    for i, ref in enumerate(references, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(14)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        r = p.add_run(f'[{i}] {ref}')
        set_run(r, size=11)

    # Save
    output_path = r'c:\Users\harin\Desktop\smart Glasses\PROJECT_SYNOPSIS.docx'
    doc.save(output_path)
    print(f'Word document generated successfully: {output_path}')


if __name__ == '__main__':
    generate_docx()
