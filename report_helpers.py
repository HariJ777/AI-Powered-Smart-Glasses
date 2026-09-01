"""Helper functions for report generation"""
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DEPT = 'Department of Computer Science(IoT Cybersecurity Including Blockchain)'
COLLEGE = "Alva's Institute of Engineering & Technology, Moodbidri"
GUIDE = 'Prof. Jyothibha R Chinchankar'
MEMBERS = [
    ('Arya B Shetty', '4AL23IC007'),
    ('Harinand J', '4AL23IC012'),
    ('Karthik P', '4AL23IC016'),
    ('Pavan SN', '4AL23IC033'),
]

def set_run(run, size=12, bold=False, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = 'Times New Roman'

def add_centered(doc, text, size=12, bold=False, after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(16)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p

def add_body(doc, text, size=12, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(18)
    r = p.add_run(text)
    set_run(r, size=size)
    return p

def add_heading1(doc, text, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.upper())
    set_run(r, size=14, bold=True)

def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, size=13, bold=True)

def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run(r, size=12, bold=True)

def add_bullet(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(18)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run('\u2022 ' + text)
    set_run(r, size=size)

def add_numbered(doc, num, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(18)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run(f'{num}. {text}')
    set_run(r, size=size)

def make_title_page(doc):
    import os
    vtu_img = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'images', 'vtu_header.png')
    if os.path.exists(vtu_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(vtu_img, width=Inches(5.5))

    add_centered(doc, 'A Mini Project Report on', size=13, after=6)
    add_centered(doc, 'AI-POWERED SMART GLASSES', size=16, bold=True, after=4)
    add_centered(doc, 'USING ESP32-CAM', size=14, bold=True, after=10)
    add_centered(doc, 'Submitted in partial fulfillment of the requirements', size=11, after=2)
    add_centered(doc, 'for the award of the degree of', size=11, after=2)
    add_centered(doc, 'BACHELOR OF ENGINEERING', size=12, bold=True, after=2)
    add_centered(doc, 'in', size=11, after=2)
    add_centered(doc, DEPT, size=11, bold=True, after=8)
    add_centered(doc, 'by', size=12, bold=True, after=6)

    table = doc.add_table(rows=len(MEMBERS), cols=2)
    for i, (name, usn) in enumerate(MEMBERS):
        c1 = table.cell(i, 0)
        p = c1.paragraphs[0]
        r = p.add_run(name)
        set_run(r, size=12)
        r.underline = True
        c2 = table.cell(i, 1)
        p = c2.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(usn)
        set_run(r, size=12)
        r.underline = True

    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        b = tblPr.find(qn('w:tblBorders'))
        if b is not None:
            tblPr.remove(b)

    add_centered(doc, '', size=6, after=8)
    add_centered(doc, 'Under the Guidance of', size=12, bold=True, after=6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(GUIDE)
    set_run(r, size=12, bold=True, italic=True)
    add_centered(doc, 'Senior Associate Professor, ' + DEPT, size=10, after=2)
    add_centered(doc, COLLEGE, size=10, after=6)
    add_centered(doc, '2025-2026', size=12, bold=True, after=2)
