"""Report content chapters 1-4"""

def write_abstract(doc, h1, h2, body, bullet):
    h1(doc, 'ABSTRACT', new_page=True)
    body(doc, 'This project presents the design and implementation of AI-powered smart glasses built using the ESP32-CAM microcontroller. The system employs a hybrid cloud-local architecture where computationally intensive AI tasks such as image description and question answering are processed through cloud-based large language models (Llama 4 Scout 17B for vision, Llama 3.3 70B for language) accessed via the Groq free-tier API, while lightweight tasks such as optical character recognition and speech-to-text transcription are handled locally using EasyOCR and OpenAI Whisper-tiny respectively.')
    body(doc, 'The backend server is built using Python Flask and implements a Smart Routing architecture that ensures only one local AI model is loaded into memory at any given time. This design decision keeps peak RAM usage under 600 MB, making the system compatible with entry-level consumer GPUs such as the NVIDIA RTX 3050. The glasses capture images through the onboard OV2640 camera module and transmit them over Wi-Fi to the backend server for processing.')
    body(doc, 'The system provides five core AI capabilities: scene description, text extraction via OCR, speech-to-text transcription, question answering, and language translation. Testing has demonstrated response times under 2 seconds for all AI tasks. The entire hardware prototype costs approximately Rs.1,407, making it a practical and affordable assistive technology solution for visually impaired individuals.')
    body(doc, 'Keywords: Smart Glasses, ESP32-CAM, Artificial Intelligence, Computer Vision, IoT, Groq API, Flask, EasyOCR, Whisper, Assistive Technology')

def write_acknowledgement(doc, h1, body, add_centered):
    h1(doc, 'ACKNOWLEDGEMENT', new_page=True)
    body(doc, 'We would like to express our sincere gratitude to our project guide, Prof. Jyothibha R Chinchankar, Senior Associate Professor, Department of Computer Science (IoT Cybersecurity Including Blockchain), for her invaluable guidance, constant encouragement, and constructive suggestions throughout the course of this project work.')
    body(doc, 'We are deeply grateful to the Head of the Department for providing us with the necessary facilities and infrastructure to carry out this project successfully. We also extend our thanks to all the faculty members of the department for their continuous support and motivation.')
    body(doc, 'We express our heartfelt thanks to the Principal and Management of Alva\'s Institute of Engineering & Technology, Moodbidri, for providing an excellent academic environment and all the resources required for the successful completion of this project.')
    body(doc, 'Finally, we would like to thank our families and friends for their unwavering support, patience, and encouragement throughout this endeavor.')
    add_centered(doc, '', size=12, after=20)
    add_centered(doc, 'Arya B Shetty (4AL23IC007)', size=12, after=2)
    add_centered(doc, 'Harinand J (4AL23IC012)', size=12, after=2)
    add_centered(doc, 'Karthik P (4AL23IC016)', size=12, after=2)
    add_centered(doc, 'Pavan SN (4AL23IC033)', size=12, after=2)

def write_toc(doc, h1, body):
    from report_helpers import add_centered
    h1(doc, 'TABLE OF CONTENTS', new_page=True)
    items = [
        ('', 'Abstract', ''),
        ('', 'Acknowledgement', ''),
        ('', 'List of Figures', ''),
        ('1', 'Introduction', ''),
        ('1.1', '   Overview', ''),
        ('1.2', '   Purpose', ''),
        ('1.3', '   Project Relevance', ''),
        ('2', 'Literature Survey', ''),
        ('3', 'Problem Statement and Objectives', ''),
        ('3.1', '   Problem Statement', ''),
        ('3.2', '   Objectives', ''),
        ('4', 'System Requirements and Specification', ''),
        ('4.1', '   Hardware Requirements', ''),
        ('4.2', '   Software Requirements', ''),
        ('4.3', '   Power Requirements', ''),
        ('5', 'Proposed Methodology', ''),
        ('5.1', '   Software Development', ''),
        ('5.2', '   Testing and Simulation', ''),
        ('6', 'System Architecture', ''),
        ('7', 'Expected Outcome', ''),
        ('8', 'Advantages and Future Scope', ''),
        ('9', 'Conclusion', ''),
        ('10', 'References', ''),
    ]
    for num, title, _ in items:
        from docx.shared import Pt, Cm
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = Pt(20)
        if num and not '.' in num:
            from report_helpers import set_run
            r = p.add_run(f'{num}.  {title}')
            set_run(r, size=12, bold=True)
        elif num:
            from report_helpers import set_run
            r = p.add_run(f'    {num}  {title}')
            set_run(r, size=12)
        else:
            from report_helpers import set_run
            r = p.add_run(f'    {title}')
            set_run(r, size=12, bold=True)

def write_ch1(doc, h1, h2, body, bullet):
    h1(doc, 'CHAPTER 1: INTRODUCTION', new_page=True)
    h2(doc, '1.1 Overview')
    body(doc, 'The rapid advancement of Artificial Intelligence and embedded systems has opened new frontiers in wearable technology. Smart glasses represent one of the most promising applications at the intersection of these technologies, offering the potential to augment human perception through real-time computational assistance. This project presents the design, development, and implementation of AI-powered smart glasses that leverage the ESP32-CAM microcontroller as the primary hardware platform.')
    body(doc, 'The ESP32-CAM is a compact, low-cost development board featuring a dual-core 32-bit processor with integrated Wi-Fi and Bluetooth connectivity, coupled with an OV2640 2-megapixel camera module. These capabilities make it an ideal candidate for building a wearable device that can capture visual information and communicate wirelessly with a backend processing server. The microcontroller captures images through its onboard camera and transmits them over a Wi-Fi network to a Python Flask backend server running on a local PC or laptop.')
    body(doc, 'The backend server implements a novel Smart Routing architecture that intelligently distributes incoming AI tasks across multiple processing pathways. Computationally intensive tasks such as image understanding and natural language question answering are offloaded to cloud-based large language models accessed through the Groq API free tier. This tier provides access to state-of-the-art models including Llama 4 Scout 17B for multimodal vision tasks and Llama 3.3 70B Versatile for language processing, with a generous allowance of 14,400 free API calls per day.')
    body(doc, 'Lightweight tasks that benefit from low latency and offline capability are processed locally using open-source AI models. EasyOCR handles optical character recognition for extracting text from images, supporting over 80 languages. OpenAI Whisper in its tiny configuration (39 million parameters) provides speech-to-text transcription capability. The Smart Router ensures that only one local model occupies memory at any given time, and models are explicitly unloaded after each request. This memory management strategy keeps peak RAM usage under 600 MB.')

    h2(doc, '1.2 Purpose')
    body(doc, 'The primary purpose of this project is to develop an affordable, functional smart glasses prototype that demonstrates the practical application of AI-powered assistive technology. The project aims to bridge the gap between expensive commercial smart glasses and the needs of users who require visual assistance, particularly visually impaired individuals in developing countries where cost is a significant barrier to adoption.')
    body(doc, 'Commercial smart glasses such as Google Glass Enterprise Edition (approximately Rs.83,000) and Meta Ray-Ban Smart Glasses (approximately Rs.25,000) are priced well beyond the reach of most users in India. By utilizing the ESP32-CAM platform and free-tier cloud AI services, this project demonstrates that comparable AI capabilities can be delivered at a fraction of the cost, with the entire hardware prototype costing under Rs.1,500.')
    body(doc, 'The project also serves as a proof-of-concept for the hybrid cloud-local AI architecture, demonstrating that sophisticated AI workloads can be efficiently distributed between cloud and edge resources without requiring expensive dedicated hardware or paid cloud subscriptions.')

    h2(doc, '1.3 Project Relevance')
    body(doc, 'This project is relevant to multiple domains within computer science and engineering:')
    bullet(doc, 'Internet of Things (IoT): The system demonstrates IoT principles through wireless sensor data collection, cloud processing, and actuator feedback via the OLED display and speaker.')
    bullet(doc, 'Embedded Systems: The ESP32-CAM firmware development involves real-time programming, peripheral management (camera, I2C, I2S, SPI), and resource-constrained optimization.')
    bullet(doc, 'Artificial Intelligence: The project integrates multiple AI paradigms including computer vision, natural language processing, optical character recognition, and automatic speech recognition.')
    bullet(doc, 'Cloud Computing: The hybrid architecture demonstrates practical cloud-edge computing patterns, API integration, and workload distribution strategies.')
    bullet(doc, 'Assistive Technology: The system directly addresses the needs of visually impaired users, aligning with the United Nations Sustainable Development Goal 10 (Reduced Inequalities).')

def write_ch2(doc, h1, h2, body, bullet):
    h1(doc, 'CHAPTER 2: LITERATURE SURVEY', new_page=True)
    body(doc, 'A comprehensive review of existing literature and related work was conducted to understand the current state of smart glasses technology and identify gaps that this project aims to address.')

    h2(doc, '2.1 Low-Cost Smart Glasses using ESP32')
    body(doc, 'Das, Saxena, and Rout (2021) presented a low-cost smart glass implementation using the ESP-32 microcontroller at the IEEE 2nd International Conference on Applied Electromagnetics. Their work demonstrated the feasibility of building wearable visual assistance devices using affordable microcontroller platforms. However, their implementation was limited to basic image capture and display without AI-powered analysis capabilities. Our project extends this approach by adding a full AI processing pipeline with cloud and local model integration.')

    h2(doc, '2.2 AI-Powered Smart Glasses for the Blind')
    body(doc, 'Ramesh Babu et al. (2025) proposed AI-powered smart glasses for visually impaired users using OpenCV and Python at the IEEE Wireless Antenna and Microwave Symposium. Their system focused on real-time object detection using traditional computer vision techniques. While effective for basic object identification, their approach lacked the ability to provide natural language descriptions of scenes or answer contextual questions about the visual environment. Our project addresses this limitation by incorporating large language models capable of generating detailed, context-aware descriptions and engaging in question-answering interactions.')

    h2(doc, '2.3 Deep Learning-Based Smart Glasses')
    body(doc, 'Lin et al. (2020) developed a smart glasses application system for visually impaired people based on deep learning at the Indo-Taiwan ICAN conference. Their work employed convolutional neural networks for object classification. However, their system required a powerful GPU-equipped computer for real-time inference, limiting its portability and accessibility. Our hybrid cloud-local architecture overcomes this constraint by offloading heavy AI tasks to cloud APIs while keeping the local compute requirements minimal.')

    h2(doc, '2.4 Speech Recognition for Wearable Devices')
    body(doc, 'Radford et al. (2023) introduced Whisper, a robust speech recognition system trained on 680,000 hours of multilingual audio data, at ICML 2023. The Whisper model family ranges from tiny (39M parameters) to large (1.5B parameters), offering a spectrum of accuracy-efficiency tradeoffs. Our project utilizes the Whisper-tiny model, which achieves acceptable accuracy for voice command recognition while requiring only approximately 145 MB of RAM, making it suitable for deployment on consumer hardware.')

    h2(doc, '2.5 Large Language Models for Visual Understanding')
    body(doc, 'Touvron et al. (2023) published the Llama 2 family of open foundation and fine-tuned chat models from Meta AI Research. These models demonstrated that open-source language models could achieve performance competitive with proprietary systems. The subsequent Llama 3.3 and Llama 4 Scout models, which our project utilizes through the Groq API, represent further improvements in both language understanding and multimodal reasoning capabilities.')

    h2(doc, '2.6 Edge Computing for AI Applications')
    body(doc, 'Chen and Ran (2019) provided a comprehensive review of deep learning with edge computing in the Proceedings of the IEEE. Their survey highlighted the importance of distributing AI workloads between cloud and edge devices to achieve optimal latency, bandwidth utilization, and privacy. Our Smart Routing architecture directly implements the principles outlined in their survey, demonstrating a practical application of cloud-edge AI distribution in a wearable device context.')

    h2(doc, '2.7 Research Gap')
    body(doc, 'The literature review reveals that while several smart glasses projects exist, most suffer from one or more of the following limitations: (1) high cost making them inaccessible to users in developing countries, (2) heavy computational requirements necessitating expensive GPU hardware, (3) limited AI capabilities restricted to basic object detection without natural language understanding, and (4) no memory optimization for running multiple AI models on consumer hardware. This project addresses all four limitations through its affordable hardware design, hybrid cloud-local architecture, multi-modal AI capabilities, and Smart Routing memory management system.')
